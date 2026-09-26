"""Generate genuine Office artifacts solely from supplied, traceable report data."""
import asyncio
from pathlib import Path
from datetime import datetime, timezone
from app.config import settings


def _exports_dir(investigation_id):
    import re
    if not re.fullmatch(r"[a-zA-Z0-9_-]{1,128}", investigation_id):
        raise ValueError("Invalid export identifier")
    base = Path(settings.SQLITE_DB_PATH).resolve().parent / "exports" / investigation_id
    base.mkdir(parents=True, exist_ok=True)
    return base


def _source_text(finding):
    return "; ".join(f"{e.get('source_id', '')} page {e.get('page', 'n/a')}" for e in finding.get("evidence", [])) or "No source recorded"


def _specialist_sections(report):
    """Keep original specialist results in exports even when synthesis omits them."""
    sections = []
    for page in report.get('visual_coverage', []):
        sections.append(('Visual page coverage',
                         f"{page['filename']} / {page['source_id']}, page {page['page']}: {page['status']}"))
    for item in report.get('visual_observations', []):
        sections.append(('Visual observation - human comparison required',
                         f"Source: {item['filename']} / {item['source_id']}, page {item['page']}\n{item['observation']}"))
    for table in report.get('calculation_results', []):
        for row in table.get('row_provenance', []):
            import json
            sections.append(('Calculation source record', json.dumps(row, ensure_ascii=False)))
        lines = [f"Source: {table['filename']} / {table['source_id']}; {table['table']}; page {table.get('page') or 'n/a'}",
                 'Calculated from original table rows by the bounded local calculator. Units follow the column headers.']
        for metric in table['results']:
            line = f"{metric['column']} - {metric['operation']}: {metric['value']}; valid rows: {metric['valid_rows']}; missing: {metric['missing_rows']}"
            if metric.get('formula'):
                line += f". Formula: {metric['formula']}; {metric['first_date']} ({metric['first_value']}) to {metric['last_date']} ({metric['last_value']}); result in percent."
            if metric.get('matching_rows'):
                line += f". Matching source rows: {metric['matching_rows']}"
            lines.append(line)
        lines.append(table['limitations'])
        sections.append(('Calculated source data - review applicability', '\n'.join(lines)))
    return sections


def _docx(investigation_id, content, report):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading("Note for review", 0)
    doc.add_paragraph(f"Investigation: {investigation_id}")
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).isoformat()}")
    doc.add_paragraph(f"Request: {report.get('query', '')}")
    doc.add_paragraph(f"Verification: {report.get('verification_status', 'unverified')}; human approval: {report.get('approval_status', 'pending')}")
    for finding in report.get('findings', []):
        doc.add_heading(str(finding.get('title', 'Finding')), 2)
        doc.add_paragraph(str(finding.get('detail', '')))
        doc.add_paragraph(f"Status: {finding.get('verification_status', 'unverified')}. Sources: {_source_text(finding)}")
    for title, detail in _specialist_sections(report):
        doc.add_heading(title, 2)
        doc.add_paragraph(detail)
    doc.add_heading("Conclusion", 2)
    doc.add_paragraph(content or report.get('conclusion') or "No conclusion supplied.")
    if report.get('recommendations'):
        doc.add_heading("Proposed actions (require review)", 2)
        doc.add_paragraph(str(report['recommendations']))
    path = _exports_dir(investigation_id) / 'note_for_approval.docx'
    doc.save(path)
    return str(path)


async def generate_briefing_docx(investigation_id, content, report=None):
    return await asyncio.to_thread(_docx, investigation_id, content, report or {})


def _xlsx(investigation_id, report):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Findings'
    ws.append(['Title', 'Detail', 'Verification', 'Sources'])
    for f in report.get('findings', []):
        ws.append([str(f.get('title', '')), str(f.get('detail', '')), str(f.get('verification_status', 'unverified')), _source_text(f)])
    if _specialist_sections(report):
        specialist = wb.create_sheet('Specialist evidence')
        specialist.append(['Analysis', 'Source and results'])
        for title, detail in _specialist_sections(report):
            specialist.append([title, detail])
            for cell in specialist[specialist.max_row]:
                cell.data_type = 's'
    telemetry = wb.create_sheet('Telemetry')
    telemetry.append(['Timestamp', 'Value', 'Unit'])
    for point in report.get('telemetry_trend') or []:
        telemetry.append([point.get('label', point.get('timestamp', '')), point.get('value'), point.get('unit', '')])
    calc = wb.create_sheet('Calculation')
    calc.append(['Measured value', 'Limit A', 'Limit B', 'Limit C', 'Assessment'])
    values = report.get('calculation_inputs') or {}
    # No measured values or engineering limits are invented. Blank inputs yield no verdict.
    calc.append([values.get('measured'), values.get('limit_a'), values.get('limit_b'), values.get('limit_c'),
                 '=IF(COUNT(A2:D2)<4,"Inputs required",IF(OR(B2>C2,C2>D2),"Invalid limits",IF(A2<=B2,"Within A",IF(A2<=C2,"Within B",IF(A2<=D2,"Within C","Above C")))))'])
    calc.append(['Limits must be supplied with their source, units and applicability.'])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]: cell.font = Font(bold=True)
        for col in ['A', 'B', 'C', 'D', 'E']: sheet.column_dimensions[col].width = 32
        for row in sheet:
            for cell in row:
                # Prevent imported text becoming executable spreadsheet formulas.
                if cell is not calc['E2'] and isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                    cell.data_type = 's'
    path = _exports_dir(investigation_id) / 'incident_log.xlsx'
    wb.save(path)
    return str(path)


async def generate_incident_xlsx(investigation_id, report=None):
    return await asyncio.to_thread(_xlsx, investigation_id, report or {})


def _pptx(investigation_id, report):
    from pptx import Presentation
    from pptx.util import Pt
    presentation = Presentation()
    sections = [('Investigation briefing', report.get('query', ''))]
    sections.extend((str(f.get('title', 'Finding')), str(f.get('detail', '')) + '\nSources: ' + _source_text(f)) for f in report.get('findings', []))
    sections.extend(_specialist_sections(report))
    sections.append(('Conclusion - pending human review', report.get('conclusion', 'No conclusion supplied.')))
    for title, text in sections:
        # Split long findings into separate slides instead of overflowing a text box.
        pieces = [str(text)[i:i+800] for i in range(0, max(len(str(text)), 1), 800)]
        for i, piece in enumerate(pieces):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title[:100] + (' (continued)' if i else '')
            frame = slide.placeholders[1].text_frame
            frame.text = piece
            for paragraph in frame.paragraphs: paragraph.font.size = Pt(18)
    path = _exports_dir(investigation_id) / 'briefing_deck.pptx'
    presentation.save(path)
    return str(path)


async def generate_briefing_pptx(investigation_id, report=None):
    return await asyncio.to_thread(_pptx, investigation_id, report or {})


async def generate_all_exports(investigation_id, report=None):
    report = report or {}
    return [
        {'format': 'docx', 'path': await generate_briefing_docx(investigation_id, report.get('conclusion', ''), report), 'label': 'Note for review'},
        {'format': 'xlsx', 'path': await generate_incident_xlsx(investigation_id, report), 'label': 'Findings and telemetry'},
        {'format': 'pptx', 'path': await generate_briefing_pptx(investigation_id, report), 'label': 'Briefing slides'},
    ]
