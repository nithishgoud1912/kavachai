"""Bounded XLSX reading. Formulas are recorded, never executed or guessed."""
import io
from datetime import date, datetime
from openpyxl import load_workbook
from app.ingestion.limits import validate_container


def read_sheets(data: bytes, filename: str) -> list[dict]:
    validate_container(data, filename)
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    sheets = []
    try:
        if len(formulas.worksheets) > 20:
            raise ValueError('Workbook exceeds 20 sheets; split the input')
        for index, sheet in enumerate(formulas.worksheets, 1):
            if sheet.max_row > 10001 or sheet.max_column > 64:
                raise ValueError(f'{filename}: worksheet exceeds 10000 data rows or 64 columns')
            rows, provenance = [], []
            for row_index, (formula_row, cached_row) in enumerate(zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1):
                values, cells = [], []
                for original, cache in zip(formula_row, cached_row):
                    formula = original.value if original.data_type == 'f' else None
                    value = cache.value if formula else original.value
                    if isinstance(value, (date, datetime)):
                        value = value.isoformat()
                    values.append(value)
                    if original.value is not None:
                        cells.append({'cell': original.coordinate, 'formula': formula,
                                      'value': value, 'cached_value_available': cache.value is not None if formula else None})
                rows.append(values)
                provenance.append({'sheet': sheet.title, 'row': row_index, 'cells': cells})
            sheets.append({'name': sheet.title, 'page': index, 'rows': rows, 'provenance': provenance})
    finally:
        formulas.close()
        cached.close()
    return sheets


def extract_spreadsheet(data: bytes, filename: str) -> list[dict]:
    import json
    pages = []
    for sheet in read_sheets(data, filename):
        for row in sheet['provenance']:
            if row['cells']:
                pages.append({'page': sheet['page'],
                              'text': f"Sheet {sheet['name']}, row {row['row']}: " + json.dumps(row['cells'], ensure_ascii=False),
                              'metadata': {'format': 'xlsx', 'sheet': sheet['name'], 'row': row['row'],
                                           'ocr_engine': 'spreadsheet', 'cells': json.dumps(row['cells'], ensure_ascii=False)}})
    return pages
