"""Mixed attachments must invoke real specialist boundaries, not relabel a text task."""
import io
import json
import uuid
from unittest.mock import AsyncMock, patch

import pymupdf
import pytest
from httpx import ASGITransport, AsyncClient
from PIL import Image
from pydantic import ValidationError

from app.services.inspection_analysis import (
    CalculationPlan, Metric, NumericTable, evaluate_plan, inspect_attachments,
)
from app.db.object_store import object_store
from app.orchestrator.model_router import model_router


def test_vision_image_is_bounded_and_transparency_is_readable():
    from app.orchestrator.model_router import prepare_vision_image
    from app.config import settings
    stream = io.BytesIO()
    Image.new('RGBA', (2048, 1024), (0, 0, 0, 0)).save(stream, format='PNG')
    with patch.object(settings, 'VISION_MAX_IMAGE_EDGE', 1024):
        result = Image.open(io.BytesIO(prepare_vision_image(stream.getvalue())))
    assert result.size == (1024, 512)
    assert result.getpixel((0, 0)) == (255, 255, 255)


def test_calculator_uses_full_rows_and_sorts_dates():
    table = NumericTable('s', 'measurements.csv', None, 'measurements', ['date', 'value'], [
        {'date':'2026-09-25', 'value':'4.6'},
        {'date':'2026-09-18', 'value':'2.0'},
        {'date':'2026-09-19', 'value':'3.0'},
    ])
    result = evaluate_plan(table, CalculationPlan(order_by='date', metrics=[
        Metric(column='value', operation='mean'), Metric(column='value', operation='percent_change')]))
    assert float(result[0]['value']) == 3.2
    assert float(result[1]['value']) == 130
    assert result[1]['first_row'] == 3
    assert result[1]['last_row'] == 2
    assert result[1]['first_date'] == '2026-09-18'


@pytest.mark.parametrize('raw', [
    '{"code":"open(\"secret\").read()","metrics":[]}',
    '{"metrics":[{"column":"x","operation":"eval"}]}',
    '{"metrics":[{"column":"x","operation":"mean","formula":"1+2"}]}',
])
def test_calculation_plan_cannot_contain_executable_instructions(raw):
    with pytest.raises(ValidationError):
        CalculationPlan.model_validate_json(raw)


@pytest.mark.parametrize('values,dates', [
    (['0', '4'], ['2026-01-01', '2026-01-02']),
    (['nan', '4', '5'], ['2026-01-01', '2026-01-02', '2026-01-03']),
    (['2', '4'], ['2026-01-01', '2026-01-01']),
    (['', '2', '4'], ['2026-01-01', '2026-01-02', '2026-01-03']),
])
def test_calculator_rejects_invalid_percent_change(values, dates):
    table = NumericTable('s', 'data.csv', None, 'data', ['date', 'value'],
                         [dict(date=date, value=value) for date, value in zip(dates, values)])
    with pytest.raises(ValueError):
        evaluate_plan(table, CalculationPlan(order_by='date', metrics=[Metric(column='value', operation='percent_change')]))


def test_long_form_telemetry_is_separated_by_asset_metric_and_unit():
    data = b'date,equipment_id,metric,value,unit\n2026-01-01,A,vibration,2,mm/s\n2026-01-02,A,vibration,4,mm/s\n2026-01-01,B,temperature,50,C\n2026-01-02,B,temperature,60,C\n'
    with patch.object(object_store, 'get_raw_file', return_value=(data, 'telemetry.csv')):
        inputs = inspect_attachments([{'source_id':'s'}])
    assert not inputs.visuals
    assert len(inputs.tables) == 2
    plan = CalculationPlan(metrics=[Metric(column='value', operation='mean')])
    assert [float(evaluate_plan(table, plan)[0]['value']) for table in inputs.tables] == [3, 55]
    assert inputs.tables[1].row_numbers == [4, 5]


def _mixed_pdf():
    image = io.BytesIO()
    Image.new('RGB', (300, 400), '#447B81').save(image, format='PNG')
    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text((50, 60), 'Inspection of test pump. The visual appendix requires review. ' * 3, fontsize=8)
        page = doc.new_page()
        page.insert_image(pymupdf.Rect(50, 50, 500, 700), stream=image.getvalue())
        return doc.tobytes()


def test_auto_detects_visual_appendix_without_reading_every_text_page():
    with patch.object(object_store, 'get_raw_file', return_value=(_mixed_pdf(), 'inspection.pdf')):
        inputs = inspect_attachments([{'source_id':'s'}])
        explicit = inspect_attachments([{'source_id':'s'}], force_vision=True)
    assert [p.page for p in inputs.visuals] == [2]
    assert [p.page for p in explicit.visuals] == [1, 2]


def test_numeric_pdf_table_is_routed_to_coder():
    with pymupdf.open() as doc:
        page = doc.new_page()
        for y in [60, 90, 120, 150]:
            page.draw_line((50, y), (350, y))
        for x in [50, 200, 350]:
            page.draw_line((x, 60), (x, 150))
        for index, (date, value) in enumerate([('Date', 'Vibration'), ('2026-01-01', '2.0'), ('2026-01-02', '4.0')]):
            page.insert_text((60, 80+30*index), date, fontsize=10)
            page.insert_text((210, 80+30*index), value, fontsize=10)
        data = doc.tobytes()
    with patch.object(object_store, 'get_raw_file', return_value=(data, 'readings.pdf')):
        inputs = inspect_attachments([{'source_id':'s'}])
    assert len(inputs.tables) == 1
    assert inputs.tables[0].page == 1
    assert inputs.tables[0].numeric_columns == ['Vibration']


@pytest.mark.asyncio
@pytest.mark.parametrize('vision_fails', [False, True])
async def test_pdf_plus_csv_uses_all_three_models_in_one_word_task(vision_fails):
    from app.main import app
    from app.db.database import init_db, async_session
    from app.db.sql_models import Session, WorkbenchJob
    from docx import Document as Word
    await init_db()
    async with async_session() as db:
        owner = Session(name='Multimodal test', department='Test')
        db.add(owner)
        await db.commit()
        owner_id = owner.id
    headers = {'Authorization': f'Bearer {owner_id}'}
    calls = []

    async def generate(**kwargs):
        calls.append(kwargs['task_type'])
        if kwargs['task_type'] == 'coding':
            return json.dumps({'order_by':'date', 'metrics':[
                {'column':'vibration', 'operation':'mean'},
                {'column':'vibration', 'operation':'percent_change'}]})
        if 'Verification Agent' in kwargs.get('system', ''):
            return json.dumps({'findings':[{'id':'f1','verification_status':'supported'}], 'overall_confidence':75})
        assert 'Model-derived visual observation' in kwargs['prompt']
        assert 'COMPUTED DATA' in kwargs['prompt']
        return json.dumps({'findings':[{'id':'f1','title':'Guard observation',
            'detail':'A guard fixing point appears empty; review the original image.', 'evidence_refs':['E1']}]})

    vision = AsyncMock(side_effect=RuntimeError('Vision model unavailable')) if vision_fails else AsyncMock(return_value='A guard fixing point appears empty.')
    with patch.object(model_router, 'embed', AsyncMock(side_effect=lambda texts:[[0.1]*768 for _ in texts])), \
         patch.object(model_router, 'generate', side_effect=generate), \
         patch.object(model_router, 'generate_vision', vision), \
         patch('app.ingestion.extract._ocr_tesseract', AsyncMock(return_value=[{'page':1, 'text':'Test inspection visual appendix. The original image requires inspection for physical condition.', 'metadata':{}}])), \
         patch('app.routers.workbench.run_code_in_sandbox', AsyncMock(side_effect=AssertionError('Document calculations must not execute arbitrary Python'))) as sandbox:
        async with AsyncClient(transport=ASGITransport(app), base_url='http://test') as client:
            attachments = []
            for filename, data, mime in [('inspection.pdf', _mixed_pdf(), 'application/pdf'),
                ('readings.csv', b'date,vibration\n2026-09-18,2.0\n2026-09-25,4.0\n', 'text/csv')]:
                upload = await client.post('/api/v1/conversations/upload', headers=headers, files={'file':(filename, data, mime)})
                assert upload.status_code == 200, upload.text
                attachments.append(dict(filename=filename, source_id=upload.json()['source_id'], url=upload.json()['url'], type='document'))
            created = await client.post('/api/v1/tasks', headers=headers, json={
                'query':'Explain these attachments and use Python-style calculations in a manager report',
                'mode':'auto', 'deliverable_type':'word', 'attachments':attachments})
            assert created.status_code == 202, created.text
            task = (await client.get('/api/v1/tasks/'+created.json()['id'], headers=headers)).json()
            assert [item['task_type'] for item in task['plan']] == ['embedding','vision','coding','text_reasoning','local_export']
            assert vision.await_count == 1
            sandbox.assert_not_awaited()
            if vision_fails:
                assert task['visual_coverage'][0]['status'] == 'failed'
                assert task['status'] == 'failed'
                assert not task['artifacts']
                assert task['plan'][1]['status'] == 'failed'
                assert task['tool_calls'][-1]['status'] == 'failed'
                assert 'coding' not in calls
                return
            assert task['status'] == 'awaiting_review', task
            assert task['visual_coverage'][0]['status'] == 'completed'
            assert set(task['models_used']) == {model_router.get_model(kind) for kind in ('embedding','vision','coding','text_reasoning')}
            assert 'coding' in calls and calls.count('text_reasoning') == 2
            assert all(item['status'] == 'done' for item in task['plan'])
            assert [item['tool_name'] for item in task['tool_calls']] == ['document_search','vision_inspect','bounded_table_calculation']
            calculated = task['calculation_results'][0]['results']
            assert float(calculated[0]['value']) == 3
            assert float(calculated[1]['value']) == 100
            file = await client.get(task['artifacts'][0]['download_url'], headers=headers)
            text = '\n'.join(p.text for p in Word(io.BytesIO(file.content)).paragraphs)
            assert 'Visual observation' in text and 'mean: 3' in text
            assert 'percent_change:' in text and '2026-09-18' in text
            async with async_session() as db:
                job = await db.get(WorkbenchJob, task['id'])
                selected = [e['data']['task_type'] for e in job.events if e['type'] == 'model_selected']
                assert selected == ['embedding','vision','coding','text_reasoning']
                assert all(e['data']['status'] == 'running' for e in job.events if e['type'] == 'tool_call_started')
