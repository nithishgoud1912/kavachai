"""P1 source regressions. Inference is mocked here; live acceptance is separate."""
import io
from unittest.mock import AsyncMock, patch
import pymupdf
import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from PIL import Image
from app.db.object_store import object_store
from app.orchestrator.model_router import model_router
from app.services.inspection_analysis import inspect_attachments, evaluate_plan, CalculationPlan, Metric


def test_generated_code_accepts_assertion_lists_but_rejects_inert_tests():
    import json
    from app.services.code_sandbox import parse_generated_code
    code, tests = parse_generated_code(json.dumps({'code': 'value = 2', 'tests': ['assert value == 2', "print('PASS')"]}))
    assert tests.startswith('assert value == 2')
    for tests in ["# assert value == 2", "def unused():\n    assert False", ['not valid Python !']]:
        with pytest.raises(ValueError):
            parse_generated_code(json.dumps({'code': code, 'tests': tests}))


def narrative(asset, value, unit, date):
    return f'Asset: {asset}\nDate: {date}\n- Pressure: {value} {unit}\n'.encode()


def test_incompatible_narrative_records_are_rejected():
    files = {'a': (narrative('A', 1, 'bar', '2026-01-01'), 'a.txt'),
             'b': (narrative('B', 100, 'kPa', '2026-01-02'), 'b.txt')}
    with patch.object(object_store, 'get_raw_file', side_effect=files.get):
        with pytest.raises(ValueError, match='No comparable'):
            inspect_attachments([{'source_id': id} for id in files])


def test_narrative_records_retain_sources_units_and_rows():
    files = {'a': (narrative('A', 1, 'bar', '2026-01-01'), 'a.txt'),
             'b': (narrative('A', 3, 'bar', '2026-01-02'), 'b.txt')}
    with patch.object(object_store, 'get_raw_file', side_effect=files.get):
        table, = inspect_attachments([{'source_id': id} for id in files]).tables
    result = evaluate_plan(table, CalculationPlan(metrics=[Metric(column='value', operation='mean')]))
    assert result[0]['value'] == '2'
    assert result[0]['unit'] == 'bar'
    assert table.source_id == 'a'
    assert [p['source_id'] for p in table.row_provenance] == ['a', 'b']
    assert all(p['unit'] == 'bar' and p['row'] == 3 and p['page'] == 1 for p in table.row_provenance)


@pytest.mark.parametrize('count', [6, 20, 21])
def test_forced_pdf_coverage_is_complete_or_rejected(count):
    with pymupdf.open() as doc:
        for _ in range(count): doc.new_page()
        data = doc.tobytes()
    with patch.object(object_store, 'get_raw_file', return_value=(data, 'scan.pdf')):
        if count > 20:
            with pytest.raises(ValueError, match='split the request'):
                inspect_attachments([{'source_id': 's'}], force_vision=True)
        else:
            assert [v.page for v in inspect_attachments([{'source_id': 's'}], force_vision=True).visuals] == list(range(1, count + 1))


async def client_and_headers():
    from app.main import app
    from app.db.database import init_db, async_session
    from app.db.sql_models import Session
    await init_db()
    async with async_session() as db:
        owner = Session(name='P1 regression', department='Test')
        db.add(owner)
        await db.commit()
        headers = {'Authorization': f'Bearer {owner.id}'}
    return AsyncClient(transport=ASGITransport(app), base_url='http://test'), headers


@pytest.mark.asyncio
@pytest.mark.parametrize('extension,fails', [('png', False), ('pdf', False), ('png', True)])
async def test_multimodal_chat_api_includes_vision_or_reports_failure(extension, fails):
    image = io.BytesIO()
    Image.new('RGB', (100, 100), 'white').save(image, format='PNG')
    data = image.getvalue()
    if extension == 'pdf':
        with pymupdf.open() as doc:
            page = doc.new_page()
            page.insert_image(page.rect, stream=data)
            data = doc.tobytes()
    client, headers = await client_and_headers()
    vision = AsyncMock(side_effect=RuntimeError('offline')) if fails else AsyncMock(return_value='UNIQUE_VISUAL_GUARD_MARKER')
    chat = AsyncMock(return_value='Observed guard marker')
    with patch.object(model_router, 'embed', AsyncMock(side_effect=lambda texts: [[0.1]*768 for _ in texts])), \
         patch('app.ingestion.extract._ocr_tesseract', AsyncMock(return_value=[{'page': 1, 'text': 'Public synthetic diagram inspection', 'metadata': {}}])), \
         patch.object(model_router, 'generate_vision', vision), patch.object(model_router, 'generate_chat', chat):
        async with client:
            upload = await client.post('/api/v1/conversations/upload', headers=headers, files={'file': (f'diagram.{extension}', data)})
            assert upload.status_code == 200, upload.text
            conversation = await client.post('/api/v1/conversations', headers=headers,
                                             json={'title': 'P1 test', 'session_id': headers['Authorization'].removeprefix('Bearer ')})
            assert conversation.status_code in (200, 201), conversation.text
            attachment = upload.json()
            response = await client.post(f"/api/v1/conversations/{conversation.json()['id']}/messages", headers=headers,
                                         json={'content': 'Describe the diagram', 'attachments': [attachment]})
            assert response.status_code == (503 if fails else 200), response.text
            if fails:
                chat.assert_not_awaited()
            else:
                messages = chat.call_args.kwargs['messages']
                assert 'UNIQUE_VISUAL_GUARD_MARKER' in messages[-1]['content']
                assert sum(m['role'] == 'user' for m in messages) == 1


@pytest.mark.asyncio
async def test_xlsx_upload_authorized_detail_and_calculation_boundary():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Readings'
    ws.append(['Date', 'Pressure_bar', 'Formula'])
    ws.append(['2026-01-01', 1, '=B2*2'])
    ws.append(['2026-01-02', 3, '=B3*2'])
    stream = io.BytesIO(); wb.save(stream)
    client, headers = await client_and_headers()
    with patch.object(model_router, 'embed', AsyncMock(side_effect=lambda texts: [[0.1]*768 for _ in texts])):
        async with client:
            upload = await client.post('/api/v1/conversations/upload', headers=headers, files={'file': ('measurements.xlsx', stream.getvalue())})
            assert upload.status_code == 200, upload.text
            source = upload.json()['source_id']
            listed = await client.get('/api/v1/knowledge-base/documents', headers=headers)
            assert any(d['document_id'] == source and d['pages'] == 1 for d in listed.json()['documents'])
            chunks = await client.get(f'/api/v1/knowledge-base/documents/{source}/chunks', headers=headers)
            assert chunks.status_code == 200
            metadata = [c['metadata'] for c in chunks.json()['chunks']]
            assert all(m['sheet'] == 'Readings' for m in metadata)
            assert any('=B2*2' in m['cells'] and 'false' in m['cells'] for m in metadata)
            from app.access import validate_attachments
            from app.db.database import async_session
            from app.db.sql_models import Session
            async with async_session() as db:
                session = await db.get(Session, headers['Authorization'].removeprefix('Bearer '))
                attached = await validate_attachments([upload.json()], session, db)
            table, = inspect_attachments(attached).tables
            assert table.rows[0]['Formula'] == ''
            assert table.row_provenance[0]['cells'][2]['formula'] == '=B2*2'
            other, other_headers = await client_and_headers()
            async with other:
                denied = await other.get(f'/api/v1/knowledge-base/documents/{source}/chunks', headers=other_headers)
                assert denied.status_code == 404
