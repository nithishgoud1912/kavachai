"""Office exports preserve input evidence and never invent process measurements."""
import uuid
import zipfile
import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from app.services.document_export import generate_all_exports,generate_incident_xlsx,generate_briefing_docx

@pytest.mark.asyncio
async def test_all_office_formats_preserve_source_and_valid_archives():
    report={'query':'Unit Z review','conclusion':'Observed fact only','verification_status':'verified','findings':[{'title':'Source finding','detail':'Observed fact only','evidence':[{'source_id':'source-123','page':3}],'verification_status':'supported'}]}
    items=await generate_all_exports(uuid.uuid4().hex,report)
    assert [i['format'] for i in items]==['docx','xlsx','pptx']
    assert all(zipfile.is_zipfile(i['path']) for i in items)
    word=' '.join(p.text for p in Document(items[0]['path']).paragraphs)
    assert 'source-123 page 3' in word and 'Observed fact only' in word
    assert 'pending' in word and 'P-102' not in word
    assert len(Presentation(items[2]['path']).slides)==3

@pytest.mark.asyncio
async def test_missing_report_never_fabricates_measurements_or_certification():
    path=await generate_incident_xlsx(uuid.uuid4().hex)
    wb=load_workbook(path)
    assert wb['Telemetry'].max_row==1 and wb['Findings'].max_row==1
    assert all(wb['Calculation'].cell(2,col).value is None for col in range(1,5))
    word=Document(await generate_briefing_docx(uuid.uuid4().hex,'Supplied note'))
    text=' '.join(p.text for p in word.paragraphs)
    for prohibited in ['ISO 10816','Zero External','P-102','95%']: assert prohibited not in text

@pytest.mark.asyncio
async def test_formula_inputs_and_injection_safety():
    path=await generate_incident_xlsx(uuid.uuid4().hex,{'calculation_inputs':{'measured':3,'limit_a':1,'limit_b':2,'limit_c':4},'findings':[{'title':'=WEBSERVICE("https://example.com")','detail':'+cmd','evidence':[]}]})
    wb=load_workbook(path)
    assert wb['Findings']['A2'].data_type=='s'
    assert wb['Calculation']['A2'].value==3
    formula=wb['Calculation']['E2'].value
    for comparison in ['A2<=B2','A2<=C2','A2<=D2','COUNT(A2:D2)<4']: assert comparison in formula

@pytest.mark.asyncio
async def test_long_presentations_paginate_instead_of_overflow():
    items=await generate_all_exports(uuid.uuid4().hex,{'findings':[{'title':'Long finding','detail':'a'*2000,'evidence':[]}]})
    presentation=Presentation(items[2]['path'])
    assert len(presentation.slides)>=5
