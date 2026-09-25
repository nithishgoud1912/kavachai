"""Regression tests for the September codebase review; no real data/model/worker."""
import asyncio
import io
import json
import uuid
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import AsyncMock, patch
import pytest
import pytest_asyncio
from httpx import AsyncClient, ASGITransport
from sqlalchemy import select
from app.main import app
from app.db.database import init_db, async_session
from app.db.sql_models import Session, Conversation, Investigation, Document, User, WorkbenchJob
from app.orchestrator.model_router import model_router
from app.agents.base import EvidenceBundle, EvidenceItem, DocumentChunk, DraftFinding, DraftFindings
from app.agents.synthesis import _resolve_evidence_refs, _format_evidence_bundle
from app.agents.base import VerificationResult, VerifiedFinding, VerificationStatus
from app.agents.verification import _build_result, _fallback_verification
from app.routers.workbench import _has_cited_partial_support
from app.services.document_export import generate_all_exports
from app.security import generate_totp_secret, generate_totp

@pytest_asyncio.fixture
async def setup():
    await init_db()
    async with async_session() as db:
        owner=Session(name='Regression owner',department='Test')
        other=Session(name='Other owner',department='Test')
        db.add_all([owner,other]); await db.commit()
        conversation=Conversation(session_id=owner.id,user='Regression',department='Test',title='PRIVATE')
        investigation=Investigation(session_id=owner.id,query='PRIVATE',status='complete',report={'conclusion':'PRIVATE','findings':[]})
        db.add_all([conversation,investigation]);await db.commit()
        return owner.id,other.id,conversation.id,investigation.id

@pytest.mark.asyncio
async def test_anonymous_and_cross_owner_reads_deletes_denied(setup):
    owner,other,conv,inv=setup
    async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
        for method,url in [('GET','/conversations'),('GET',f'/conversations/{conv}'),('DELETE',f'/conversations/{conv}'),('GET',f'/investigations/{inv}/report'),('GET','/tasks')]:
            assert (await client.request(method,'/api/v1'+url)).status_code==401
        for url in [f'/conversations/{conv}',f'/investigations/{inv}/report',f'/exports/{inv}/docx']:
            assert (await client.get('/api/v1'+url,headers={'Authorization':f'Bearer {other}'})).status_code==404
        assert (await client.get(f'/api/v1/conversations/{conv}',headers={'Authorization':f'Bearer {owner}'})).status_code==200

@pytest.mark.asyncio
async def test_expired_and_completed_mfa_challenges_rejected():
    await init_db()
    secret=generate_totp_secret()
    async with async_session() as db:
        user=User(username=uuid.uuid4().hex,password_hash='unused',department='Test',mfa_secret=secret,mfa_enabled=True)
        db.add(user);await db.flush()
        sessions=[Session(name='MFA',department='Test',user_id=user.id,mfa_verified=False,expires_at=datetime.now(timezone.utc)-timedelta(seconds=1)),
                  Session(name='MFA',department='Test',user_id=user.id,mfa_verified=True,expires_at=datetime.now(timezone.utc)+timedelta(hours=1))]
        db.add_all(sessions);await db.commit()
    async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
        for session in sessions:
            assert (await client.post('/api/v1/auth/mfa/verify',json={'session_id':session.id,'code':generate_totp(secret)})).status_code==401

@pytest.mark.asyncio
async def test_owner_retains_access_after_new_login():
    from app.access import assert_owner
    await init_db()
    async with async_session() as db:
        user=User(username=uuid.uuid4().hex,password_hash='unused',department='Test')
        db.add(user);await db.flush()
        old=Session(name='Owner',department='Test',user_id=user.id,is_revoked=True)
        fresh=Session(name='Owner',department='Test',user_id=user.id)
        db.add_all([old,fresh]);await db.flush()
        doc=Document(filename='test.txt',source_id=uuid.uuid4().hex,document_type='manual',session_id=old.id)
        db.add(doc);await db.commit()
        await assert_owner(doc,fresh,db)

@pytest.mark.asyncio
async def test_model_outage_never_returns_generated_evidence_or_vectors():
    with patch.object(model_router,'is_ollama_available',AsyncMock(return_value=False)):
        for call in [lambda:model_router.generate('summarize'),lambda:model_router.generate_chat([]),lambda:model_router.embed(['test']),lambda:model_router.generate_vision('read',b'image')]:
            with pytest.raises(RuntimeError): await call()


def test_citations_match_identity_and_missing_refs_are_not_supported():
    bundle=EvidenceBundle(evidence_items=[EvidenceItem(type='document',source_id='A',label='A'),EvidenceItem(type='document',source_id='B',label='B')])
    assert [e.source_id for e in _resolve_evidence_refs(['B'],bundle)]==['B']
    assert [e.source_id for e in _resolve_evidence_refs(['E2'],bundle)]==['B']
    assert _resolve_evidence_refs(['unknown'],bundle)==[]
    draft=DraftFindings(findings=[DraftFinding(id='f1',title='Claim',detail='Claim',evidence=[])],condition_summary='')
    assert _build_result(draft,{'findings':[{'id':'f1','verification_status':'supported'}]}).overall_status=='unverified'
    assert _fallback_verification(draft).overall_status=='unverified'

def test_synthesis_formats_short_citation_ids_for_document_chunks():
    bundle=EvidenceBundle(
        document_findings=[DocumentChunk(chunk_text='Recorded observation',source_id='opaque-source-id',page=2,score=0.9)],
        evidence_items=[EvidenceItem(type='document',source_id='opaque-source-id',label='inspection.pdf',page=2)],
    )
    assert '[E1] Source opaque-source-id (p.2)' in _format_evidence_bundle(bundle)

def test_partially_supported_findings_need_citations_to_continue_to_review():
    citation=EvidenceItem(type='document',source_id='source-1',label='inspection.pdf')
    partial=VerificationResult(findings=[VerifiedFinding(
        id='f1',title='Observation',detail='Uncertain observation',
        verification_status=VerificationStatus.PARTIALLY_SUPPORTED,evidence=[citation],
    )],overall_confidence=20,overall_status='unverified')
    unsupported=VerificationResult(findings=[VerifiedFinding(
        id='f1',title='Claim',detail='Unsupported claim',
        verification_status=VerificationStatus.UNSUPPORTED,evidence=[],
    )],overall_confidence=0,overall_status='unverified')
    assert _has_cited_partial_support(partial)
    assert not _has_cited_partial_support(unsupported)

@pytest.mark.asyncio
async def test_office_artifacts_are_valid_and_do_not_invent_telemetry():
    from docx import Document as Word
    from openpyxl import load_workbook
    from pptx import Presentation
    report={'query':'Inspect UNIT-X','conclusion':'Only supplied evidence.','findings':[{'title':'Recorded finding','detail':'UNIT-X inspected','verification_status':'supported','evidence':[{'source_id':'source-x','page':2}]}]}
    artifacts=await generate_all_exports(uuid.uuid4().hex,report)
    for artifact in artifacts: assert zipfile.is_zipfile(artifact['path'])
    text=' '.join(p.text for p in Word(artifacts[0]['path']).paragraphs)
    assert 'source-x' in text and 'P-102' not in text and 'Zero External' not in text
    wb=load_workbook(artifacts[1]['path'])
    assert wb['Telemetry'].max_row==1
    assert wb['Calculation']['A2'].value is None
    assert 'A2<=B2' in wb['Calculation']['E2'].value
    assert len(Presentation(artifacts[2]['path']).slides)==3

@pytest.mark.asyncio
async def test_task_outage_is_persisted_as_failed(setup):
    owner,*_=setup
    with patch.object(model_router,'generate',AsyncMock(side_effect=RuntimeError('model offline'))):
        async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
            response=await client.post('/api/v1/tasks',headers={'Authorization':f'Bearer {owner}'},json={'query':'write Python code','mode':'code'})
            assert response.status_code==202
            id=response.json()['id']
            task=(await client.get(f'/api/v1/tasks/{id}',headers={'Authorization':f'Bearer {owner}'})).json()
            assert task['status']=='failed' and task['artifacts']==[]

@pytest.mark.asyncio
async def test_sandbox_route_preserves_real_failure(setup):
    from app.services.code_sandbox import SandboxResult
    owner,*_=setup
    with patch('app.routers.runtime.run_code_in_sandbox',AsyncMock(return_value=SandboxResult('', 'Expected failure', 1))):
        async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
            response=await client.post('/api/v1/sandbox/run',headers={'Authorization':f'Bearer {owner}'},json={'code':'raise RuntimeError()'})
            assert response.status_code==200
            assert response.json()['status']=='runtime_error' and response.json()['exit_code']==1

@pytest.mark.asyncio
async def test_mixed_pdf_extracts_every_page():
    import pymupdf
    from app.ingestion.extract import extract_text_with_ocr
    pdf=pymupdf.open()
    pdf.new_page().insert_text((50,50),'Embedded text document content ' * 4)
    pdf.new_page()
    data=pdf.tobytes();pdf.close()
    with patch('app.ingestion.extract._ocr_tesseract',AsyncMock(return_value=[{'page':1,'text':'Scanned page fact','metadata':{}}])):
        pages=await extract_text_with_ocr(data,'mixed.pdf')
    assert [p['page'] for p in pages]==[1,2]
    assert pages[1]['text']=='Scanned page fact'


def test_object_store_and_numerical_inputs_fail_closed():
    from app.db.object_store import object_store
    from app.agents.data_agent import _compute_pct_change
    from app.ingestion.tabular import parse_tabular_file
    with pytest.raises(ValueError): object_store.get_raw_file('../outside')
    assert _compute_pct_change([0,5]) is None
    bad=b'timestamp,equipment_id,metric,value,unit\n2026-01-01,P-1,vibration,inf,mm/s\n'
    with pytest.raises(ValueError): parse_tabular_file(bad,'bad.csv',uuid.uuid4().hex)

@pytest.mark.asyncio
async def test_network_report_never_certifies_whole_deployment(setup):
    owner,*_=setup
    async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
        result=await client.get('/api/v1/network-monitor',headers={'Authorization':f'Bearer {owner}'})
        assert result.status_code==200
        assert result.json()['verified_airgap'] is False
        assert 'Python backend' in result.json()['scope']


@pytest.mark.asyncio
async def test_authenticated_upload_to_real_word_deliverable(setup):
    owner,other,*_=setup
    headers={'Authorization':f'Bearer {owner}'}
    with patch.object(model_router,'embed',AsyncMock(side_effect=lambda texts:[[0.1]*768 for _ in texts])):
        async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
            upload=await client.post('/api/v1/conversations/upload',headers=headers,files={'file':('inspection.txt',b'Unit X was inspected on 20 September. The inspection recorded a loose protective cover.','text/plain')})
            assert upload.status_code==200,upload.text
            source=upload.json()['source_id']
            synthesis={'condition_summary':'Inspection observation','findings':[{'id':'f1','title':'Loose cover','detail':'The inspection recorded a loose protective cover.','evidence_refs':[source]}]}
            verification={'findings':[{'id':'f1','verification_status':'supported'}],'overall_confidence':80}
            with patch.object(model_router,'generate',AsyncMock(side_effect=[json.dumps(synthesis),json.dumps(verification)])):
                created=await client.post('/api/v1/tasks',headers=headers,json={'query':'Draft a note from this inspection','mode':'document','deliverable_type':'word','attachments':[{'filename':'inspection.txt','url':upload.json()['url'],'source_id':source,'type':'document'}]})
            assert created.status_code==202,created.text
            task=(await client.get('/api/v1/tasks/'+created.json()['id'],headers=headers)).json()
            assert task['status']=='awaiting_review',task
            artifact=task['artifacts'][0]
            file=await client.get(artifact['download_url'],headers=headers)
            assert file.status_code==200 and zipfile.is_zipfile(io.BytesIO(file.content))
            from docx import Document as Word
            assert 'loose protective cover' in ' '.join(p.text for p in Word(io.BytesIO(file.content)).paragraphs)
            assert (await client.get(artifact['download_url'],headers={'Authorization':f'Bearer {other}'})).status_code==404
            denied=await client.patch('/api/v1/tasks/'+task['id'],headers=headers,json={'action':'approve','version':1})
            assert denied.status_code==403

@pytest.mark.asyncio
async def test_code_task_records_failed_attempt_and_successful_repair(setup):
    from app.services.code_sandbox import SandboxResult
    owner,*_=setup
    generated=[json.dumps({'code':'value=1','tests':'assert value==2'}),json.dumps({'code':'value=2','tests':'assert value==2'})]
    with patch.object(model_router,'generate',AsyncMock(side_effect=generated)), patch('app.services.code_sandbox.run_code_in_sandbox',AsyncMock()):
        with patch('app.routers.workbench.run_code_in_sandbox',AsyncMock(side_effect=[SandboxResult('','AssertionError',1),SandboxResult('passed','',0)])):
            async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
                headers={'Authorization':f'Bearer {owner}'}
                created=await client.post('/api/v1/tasks',headers=headers,json={'query':'Return 2','mode':'code'})
                task=(await client.get('/api/v1/tasks/'+created.json()['id'],headers=headers)).json()
                assert task['status']=='awaiting_review',task
                assert [c['status'] for c in task['tool_calls']]==['failed','completed']
                artifact=await client.get(task['artifacts'][0]['download_url'],headers=headers)
                assert 'value=2' in artifact.text and 'assert value==2' in artifact.text

@pytest.mark.asyncio
async def test_vision_task_reads_real_uploaded_image_and_records_both_models(setup):
    from PIL import Image
    owner,*_=setup
    image=io.BytesIO(); Image.new('RGB',(80,40),'white').save(image,format='PNG')
    headers={'Authorization':f'Bearer {owner}'}
    with patch.object(model_router,'embed',AsyncMock(side_effect=lambda texts:[[0.1]*768 for _ in texts])), patch('app.ingestion.extract._ocr_tesseract',AsyncMock(return_value=[{'page':1,'text':'Inspection cover','metadata':{}}])):
        async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
            upload=await client.post('/api/v1/conversations/upload',headers=headers,files={'file':('image.png',image.getvalue(),'image/png')})
            assert upload.status_code==200,upload.text
            source=upload.json()['source_id']
            draft={'findings':[{'id':'v1','title':'Cover','detail':'A cover is visible.','evidence_refs':[source]}]}
            checked={'findings':[{'id':'v1','verification_status':'supported'}],'overall_confidence':60}
            with patch.object(model_router,'generate_vision',AsyncMock(return_value='A cover is visible.')) as vision, patch.object(model_router,'generate',AsyncMock(side_effect=[json.dumps(draft),json.dumps(checked)])):
                result=await client.post('/api/v1/tasks',headers=headers,json={'query':'Describe the cover','mode':'vision','attachments':[{'filename':'image.png','source_id':source,'type':'image','url':upload.json()['url']}]})
                task=(await client.get('/api/v1/tasks/'+result.json()['id'],headers=headers)).json()
                assert task['status']=='awaiting_review',task
                assert vision.await_args.kwargs['image_bytes']==image.getvalue()
                assert model_router.get_model('vision') in task['models_used']
                assert all(step['status']=='done' and step.get('started_at') for step in task['plan'])

@pytest.mark.asyncio
async def test_equipment_graph_is_persistent_and_scoped(setup):
    from app.routers.knowledge_base import add_graph_node, get_graph
    from app.models.ingestion import GraphNodeCreate
    owner,other,*_=setup
    async with async_session() as db:
        session=await db.get(Session,owner)
        await add_graph_node(GraphNodeCreate(equipment_id='TEST-900'),session,None,db)
    async with async_session() as db:
        session=await db.get(Session,owner)
        assert [n['id'] for n in (await get_graph(session,None,db)).nodes]==['TEST-900']
        stranger=await db.get(Session,other)
        assert (await get_graph(stranger,None,db)).nodes==[]

@pytest.mark.asyncio
async def test_bearer_tokens_are_not_accepted_in_query_strings(setup):
    owner,*_=setup
    async with AsyncClient(transport=ASGITransport(app),base_url='http://test') as client:
        assert (await client.get('/api/v1/tasks',params={'session_id':owner})).status_code==401
